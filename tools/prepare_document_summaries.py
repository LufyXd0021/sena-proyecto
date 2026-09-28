"""Analyze project documents, create the Word PDF, and publish concise summaries."""

from pathlib import Path
import html
import shutil
import subprocess
import tempfile
from io import BytesIO

import pandas as pd
from docx import Document as WordDocument
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import Image as PdfImage, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle, PageBreak
from PIL import Image as PillowImage

import os
import sys

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'config.settings')

import django

django.setup()

from portal.models import Document


def clean(value):
    return '' if value is None else str(value).strip()


def word_summary(path):
    document = WordDocument(path)
    paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    return (
        'Informe de analítica de datos que estudia 332.390 casos y 495.272 víctimas '
        'de lesiones personales en Colombia entre 2021 y 2025. Presenta el problema, '
        'la metodología y el marco teórico, y analiza tendencias anuales y mensuales, '
        'medios empleados y género, grupos de edad, días de la semana y distribución '
        'regional. Incluye tablas de resultados, una propuesta de dashboard, estrategias '
        'de comunicación y conclusiones. Contiene '
        f'{len(paragraphs)} párrafos y {len(document.tables)} tablas.'
    )


def excel_summary(path):
    book = pd.ExcelFile(path)
    sheets = []
    total_rows = 0
    for sheet in book.sheet_names:
        frame = pd.read_excel(path, sheet_name=sheet, header=None)
        sheets.append(sheet)
        total_rows += len(frame)
    sheet_text = ', '.join(sheets[:5])
    if len(sheets) > 5:
        sheet_text += f' y {len(sheets) - 5} hojas más'
    return (
        f'Base de datos operativa del proyecto con {len(sheets)} hojas y aproximadamente '
        f'{total_rows:,} filas de información. Organiza los registros de lesiones personales '
        f'por cortes anuales y temáticos para permitir análisis temporal, demográfico y '
        f'territorial. Entre sus hojas se encuentran: {sheet_text}. Es la fuente tabular '
        'para construir indicadores, segmentaciones y visualizaciones del informe.'
    ).replace(',', '.')


def reportlab_pdf(source, destination):
    document = WordDocument(source)
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name='DocTitle', parent=styles['Title'], alignment=TA_CENTER, fontSize=16, leading=20, spaceAfter=14))
    styles.add(ParagraphStyle(name='DocBody', parent=styles['BodyText'], fontSize=9.5, leading=14, spaceAfter=7))
    story = []
    image_parts = {}
    for relationship in document.part.rels.values():
        if 'image' in relationship.reltype:
            image_parts[relationship.rId] = relationship.target_part.blob

    for paragraph in document.paragraphs:
        text = clean(paragraph.text)
        if text:
            story.append(Paragraph(html.escape(text), styles['DocTitle'] if not story else styles['DocBody']))
        for blip in paragraph._p.xpath('.//a:blip'):
            blob = image_parts.get(blip.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed'))
            if blob:
                with PillowImage.open(BytesIO(blob)) as image:
                    width, height = image.size
                max_width, max_height = 175 * mm, 220 * mm
                scale = min(max_width / width, max_height / height, 1)
                story.extend([Spacer(1, 6), PdfImage(BytesIO(blob), width=width * scale, height=height * scale), Spacer(1, 8)])
    for table in document.tables:
        rows = [[clean(cell.text).replace('\n', ' / ') for cell in row.cells] for row in table.rows]
        if rows:
            rendered = Table(rows, repeatRows=1, hAlign='LEFT')
            rendered.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#c9ed57')),
                ('GRID', (0, 0), (-1, -1), .35, colors.HexColor('#9da69d')),
                ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                ('FONTSIZE', (0, 0), (-1, -1), 7),
                ('VALIGN', (0, 0), (-1, -1), 'TOP'),
                ('LEFTPADDING', (0, 0), (-1, -1), 5),
                ('RIGHTPADDING', (0, 0), (-1, -1), 5),
            ]))
            story.extend([Spacer(1, 6), rendered, PageBreak()])
    destination.parent.mkdir(parents=True, exist_ok=True)
    SimpleDocTemplate(str(destination), pagesize=A4, rightMargin=16 * mm, leftMargin=16 * mm, topMargin=15 * mm, bottomMargin=15 * mm).build(story)


def convert_word_to_pdf(source, destination):
    destination.parent.mkdir(parents=True, exist_ok=True)
    soffice = shutil.which('soffice') or shutil.which('libreoffice')
    if soffice:
        with tempfile.TemporaryDirectory() as temp_dir:
            subprocess.run([soffice, '--headless', '--convert-to', 'pdf', '--outdir', temp_dir, str(source)], check=True, capture_output=True)
            generated = Path(temp_dir) / f'{source.stem}.pdf'
            if generated.exists():
                shutil.copyfile(generated, destination)
                return 'LibreOffice'
    reportlab_pdf(source, destination)
    return 'ReportLab'


def main():
    media = ROOT / 'media' / 'documents'
    word = next(media.glob('*.docx'))
    excel = next(media.glob('*.xlsx'))
    pdf = media / 'pdf' / f'{word.stem}.pdf'
    converter = convert_word_to_pdf(word, pdf)
    summaries = {
        word.name: word_summary(word),
        excel.name: excel_summary(excel),
    }
    word_record = Document.objects.filter(file__icontains=word.name).first()
    if word_record:
        word_record.summary = summaries[word.name]
        word_record.pdf_file.name = str(pdf.relative_to(ROOT / 'media')).replace('\\', '/')
        word_record.save(update_fields=['summary', 'pdf_file'])
    excel_record = Document.objects.filter(file__icontains=excel.name).first()
    if excel_record:
        excel_record.summary = summaries[excel.name]
        excel_record.save(update_fields=['summary'])
    report = ROOT / 'media' / 'document_summaries.txt'
    report.write_text('\n\n'.join(f'{name}\n{summary}' for name, summary in summaries.items()), encoding='utf-8')
    print(f'PDF: {pdf} ({converter})')
    print(report.read_text(encoding='utf-8'))


if __name__ == '__main__':
    main()
from django.shortcuts import get_object_or_404, redirect, render
from django.conf import settings
from pathlib import Path
import mimetypes
import base64
from collections import defaultdict
import html
import json
import socket
import xml.etree.ElementTree as ElementTree
from io import BytesIO
from zipfile import BadZipFile

from docx import Document as WordDocument
from docx.opc.exceptions import PackageNotFoundError
from docx.table import Table as WordTable
from docx.text.paragraph import Paragraph as WordParagraph
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import Image as PdfImage, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from PIL import Image as PillowImage, ImageDraw
try:
	import qrcode
except ImportError:
	qrcode = None
from django.core.files.storage import default_storage
from django.db.models import Q
from django.http import FileResponse, Http404, HttpResponse, JsonResponse
from django.views.decorators.cache import cache_page
from openpyxl import load_workbook
from pptx import Presentation

from .models import Document
from .services import build_homepage_summary, get_answer_for_question, get_chat_categories, get_quick_questions, get_question_bank
from .utils import dashboard_number as _dashboard_number, dashboard_display as _dashboard_display


def home(request):
	documents = Document.objects.filter(is_published=True)
	statistics = dashboard_preview_data()
	return render(request, 'portal/home.html', {'documents': documents, 'statistics': statistics, 'insights': homepage_insights(statistics), 'show_map': False, 'project_url': project_access_url(request)})


def map_page(request):
	return redirect('portal:territory')


@cache_page(60 * 15)
def dashboard_page(request):
	statistics = dashboard_preview_data()
	return render(request, 'portal/dashboard.html', {'statistics': statistics, 'insights': homepage_insights(statistics), 'report_summary': build_report_summary(statistics), 'project_url': project_access_url(request)})


@cache_page(60 * 5)
def dashboard_api(request):
	return JsonResponse(dashboard_preview_data())


def generated_report_pdf(request):
	if request.method != 'POST':
		return JsonResponse({'error': 'Usa POST para generar el informe.'}, status=405)
	try:
		payload = json.loads(request.body or '{}')
		report_text = str(payload.get('report_text', '')).strip()
		chart_images = payload.get('chart_images', {}) or {}
	except json.JSONDecodeError:
		return JsonResponse({'error': 'El informe no tiene un formato válido.'}, status=400)
	if not report_text:
		return JsonResponse({'error': 'No hay contenido para generar el informe.'}, status=400)

	buffer = BytesIO()
	doc = SimpleDocTemplate(buffer, pagesize=A4, rightMargin=22 * mm, leftMargin=22 * mm, topMargin=20 * mm, bottomMargin=20 * mm)
	styles = getSampleStyleSheet()
	title_style = ParagraphStyle('ReportTitle', parent=styles['Title'], fontName='Helvetica-Bold', fontSize=18, leading=23, spaceAfter=14, textColor=colors.HexColor('#17211f'))
	body_style = ParagraphStyle('ReportBody', parent=styles['BodyText'], fontName='Helvetica', fontSize=10, leading=14, spaceAfter=6, textColor=colors.HexColor('#17211f'))
	story = []
	for index, line in enumerate(report_text.splitlines()):
		text = html.escape(line).replace('  ', '&nbsp; ')
		if index == 0:
			story.append(Paragraph(text, title_style))
		elif not text:
			story.append(Spacer(1, 7))
		else:
			story.append(Paragraph(text, body_style))
	for label, data_url in chart_images.items():
		if not isinstance(data_url, str) or ',' not in data_url:
			continue
		try:
			image_bytes = base64.b64decode(data_url.split(',', 1)[1])
			story.append(Spacer(1, 10))
			story.append(Paragraph(html.escape(label), styles['Heading3']))
			story.append(PdfImage(BytesIO(image_bytes), width=160 * mm, height=82 * mm))
		except (ValueError, TypeError):
			continue
	doc.build(story)
	buffer.seek(0)
	response = HttpResponse(buffer.getvalue(), content_type='application/pdf')
	response['Content-Disposition'] = 'attachment; filename="informe-dashboard.pdf"'
	return response


def findings_page(request):
	statistics = dashboard_preview_data()
	documents = Document.objects.filter(is_published=True)
	return render(request, 'portal/findings.html', {'statistics': statistics, 'insights': homepage_insights(statistics), 'documents': documents, 'project_url': project_access_url(request)})


def methodology_page(request):
	return render(request, 'portal/methodology.html', {'project_url': project_access_url(request)})


@cache_page(60 * 15)
def report_page(request):
	statistics = dashboard_preview_data()
	return render(request, 'portal/report.html', {'statistics': statistics, 'insights': homepage_insights(statistics), 'report_summary': build_report_summary(statistics), 'project_url': project_access_url(request)})


def territory_page(request):
	return render(request, 'portal/territory.html', {'project_url': project_access_url(request)})


def homepage_insights(statistics):
	return build_homepage_summary(statistics)


def build_report_summary(statistics):
	def top_item(key, fallback_label='No disponible'):
		items = statistics.get(key, []) or []
		return max(items, key=lambda item: _dashboard_number(item.get('value')), default={'label': fallback_label, 'value': '—'})

	days = statistics.get('days', []) or []
	weekend = sum(_dashboard_number(item.get('value')) for item in days if str(item.get('label', '')).upper() in {'SABADO', 'SÁBADO', 'DOMINGO'})
	day_total = sum(_dashboard_number(item.get('value')) for item in days)
	return {
		'top_weapon': top_item('weapons', 'No disponible'),
		'top_gender': top_item('gender', 'No disponible'),
		'top_age': top_item('ages', 'No disponible'),
		'top_day': top_item('days', 'No disponible'),
		'weekend_share': f'{(weekend / day_total * 100):.1f}%'.replace('.', ',') if day_total else '—',
	}


def _chat_sources():
	return list(Document.objects.filter(is_published=True).values_list('title', flat=True))


def chat(request):
	bank = get_question_bank()
	if request.method == 'GET':
		# El banco de preguntas es estático: se cachea 1 hora para evitar
		# recalcular en cada visita a la página.
		response = JsonResponse({
			'categories': get_chat_categories(),
			'quick_questions': get_quick_questions(),
		})
		response['Cache-Control'] = 'public, max-age=3600'
		return response
	if request.method != 'POST':
		return JsonResponse({'error': 'Usa GET para cargar preguntas o POST para responder.'}, status=405)
	try:
		payload = json.loads(request.body or '{}')
	except json.JSONDecodeError:
		return JsonResponse({'error': 'La solicitud no tiene un formato válido.'}, status=400)
	category = bank.get(str(payload.get('category', '')))
	question_id = str(payload.get('question', ''))
	question = next((item for item in category['questions'] if item['id'] == question_id), None) if category else None
	if not question:
		return JsonResponse({'error': 'Selecciona una categoría y una pregunta válidas.'}, status=400)
	return JsonResponse({'answer': question['answer'], 'sources': _chat_sources(), 'category': category['label'], 'question': question['label']})


def project_access_url(request):
	if settings.PROJECT_PUBLIC_URL:
		return f'{settings.PROJECT_PUBLIC_URL}/'
	host = request.get_host().split(':', 1)[0]
	if host in {'127.0.0.1', 'localhost', '0.0.0.0'}:
		connection = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
		try:
			connection.connect(('8.8.8.8', 80))
			host = connection.getsockname()[0]
		except OSError:
			# No se pudo detectar la IP de red local; devolver localhost como fallback.
			# Configura PROJECT_PUBLIC_URL en .env para fijar la dirección correcta.
			host = '127.0.0.1'
		finally:
			connection.close()
	return f'http://{host}:{request.get_port()}/'


def project_qr(request):
	if qrcode is None:
		return HttpResponse('La dependencia qrcode no está instalada.', status=503, content_type='text/plain')
	qr = qrcode.make(project_access_url(request))
	output = BytesIO()
	qr.save(output, format='PNG')
	response = HttpResponse(output.getvalue(), content_type='image/png')
	response['Cache-Control'] = 'no-store, no-cache, must-revalidate'
	return response


def documentation_statistics():
	word_document = Document.objects.filter(document_type='word', is_published=True).first()
	if not word_document or not word_document.file:
		return {'annual': []}
	if Path(word_document.file.name).suffix.lower() == '.pdf':
		return {'annual': [], 'monthly': [], 'regions': []}
	try:
		with word_document.file.open('rb') as source:
			word = WordDocument(source)
		result = {'annual': [], 'monthly': [], 'regions': []}
		for index, key in ((0, 'annual'), (1, 'monthly'), (5, 'regions')):
			if len(word.tables) <= index:
				continue
			for row in word.tables[index].rows[1:]:
				values = [cell.text.strip().replace('\n', ' ') for cell in row.cells]
				if values and values[0]:
					result[key].append({'label': values[0], 'value': values[1] if len(values) > 1 else '', 'detail': values[2] if len(values) > 2 else ''})
		return result
	except FileNotFoundError:
		return {'annual': []}
	except (OSError, ValueError, IndexError, PackageNotFoundError, BadZipFile, KeyError, TypeError, AttributeError):
		return {'annual': []}


def document_detail(request, slug):
	document = get_object_or_404(Document, slug=slug, is_published=True)
	preview = build_preview(document)
	return render(request, 'portal/document_detail.html', {'document': document, 'preview': preview})


def serve_document_file(request, file_path):
	if not Document.objects.filter(is_published=True).filter(
		Q(file=file_path) | Q(pdf_file=file_path)
	).exists():
		raise Http404('No se encontró el documento publicado.')
	try:
		file_handle = default_storage.open(file_path, 'rb')
	except FileNotFoundError as error:
		raise Http404('El archivo del documento no está disponible.') from error
	content_type, _ = mimetypes.guess_type(file_path)
	return FileResponse(
		file_handle,
		content_type=content_type or 'application/octet-stream',
		as_attachment=Path(file_path).suffix.lower() != '.pdf',
		filename=Path(file_path).name,
	)


CHART_NS = {'c': 'http://schemas.openxmlformats.org/drawingml/2006/chart'}
DRAWING_NS = {'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'}
_dashboard_cache = {}


def chart_text(node):
	if node is None:
		return ''
	value = node.find('.//c:v', CHART_NS)
	if value is not None and value.text:
		return value.text
	rich_text = node.find('.//a:t', DRAWING_NS)
	return rich_text.text if rich_text is not None and rich_text.text else ''


def render_docx_chart(chart_blob):
	root = ElementTree.fromstring(chart_blob)
	chart_type = next((name for name in ('barChart', 'lineChart', 'pieChart', 'doughnutChart', 'areaChart') if root.find(f'.//c:{name}', CHART_NS) is not None), 'barChart')
	title = chart_text(root.find('.//c:title', CHART_NS)) or 'Gráfico del informe'
	series = []
	for serie in root.findall('.//c:ser', CHART_NS):
		label_node = serie.find('.//c:tx', CHART_NS)
		category_nodes = serie.findall('.//c:cat//c:pt/c:v', CHART_NS)
		value_nodes = serie.findall('.//c:val//c:pt/c:v', CHART_NS)
		if value_nodes:
			try:
				values = [float(node.text.replace(',', '.')) for node in value_nodes]
			except (AttributeError, ValueError):
				continue
			series.append((chart_text(label_node) or 'Serie', [node.text for node in category_nodes], values))
	if not series:
		return None
	canvas = PillowImage.new('RGB', (1200, 650), 'white')
	draw = ImageDraw.Draw(canvas)
	draw.text((55, 25), title[:90], fill='#17211f')
	left, top, right, bottom = 100, 90, 1130, 545
	maximum = max(value for _, _, values in series for value in values) or 1
	draw.line((left, bottom, right, bottom), fill='#17211f', width=2)
	draw.line((left, top, left, bottom), fill='#17211f', width=2)
	colors_for_series = ['#7e9721', '#ff836b', '#17211f', '#6e9da5']
	count = max(len(values) for _, _, values in series)
	for series_index, (label, categories, values) in enumerate(series):
		color = colors_for_series[series_index % len(colors_for_series)]
		if chart_type == 'lineChart':
			points = []
			for index, value in enumerate(values):
				x = left + ((index + .5) * (right - left) / max(count, 1))
				y = bottom - (value / maximum) * (bottom - top)
				points.append((x, y))
				draw.ellipse((x - 5, y - 5, x + 5, y + 5), fill=color)
				draw.text((x - 18, y - 24), format(value, ',.0f').replace(',', '.'), fill='#17211f')
			if len(points) > 1:
				draw.line(points, fill=color, width=4)
		else:
			bar_width = max(12, (right - left) / max(count * len(series), 1) - 8)
			for index, value in enumerate(values):
				x = left + (index * len(series) + series_index) * ((right - left) / max(count * len(series), 1)) + 8
				y = bottom - (value / maximum) * (bottom - top)
				draw.rectangle((x, y, x + bar_width, bottom), fill=color)
				draw.text((x, max(top, y - 20)), format(value, ',.0f').replace(',', '.'), fill='#17211f')
		for index, category in enumerate(categories[:count]):
			x = left + ((index + .5) * (right - left) / max(count, 1))
			draw.text((x - 18, bottom + 15), str(category)[:12], fill='#65716d')
		draw.text((right - 160, 25 + series_index * 22), f'{label[:18]}  ■', fill=color)
	return canvas


def document_pdf(request, slug):
	document = get_object_or_404(Document, slug=slug, is_published=True, document_type='word')
	download = request.GET.get('download') == '1'
	if document.pdf_file:
		return FileResponse(
			document.pdf_file.open('rb'),
			content_type='application/pdf',
			as_attachment=download,
			filename=f'{Path(document.pdf_file.name).stem}.pdf',
		)
	if not document.file:
		return HttpResponse('No se encontró el archivo del documento.', status=404)
	if Path(document.file.name).suffix.lower() == '.pdf':
		return FileResponse(
			document.file.open('rb'),
			content_type='application/pdf',
			as_attachment=download,
			filename=f'{Path(document.file.name).stem}.pdf',
		)
	try:
		with document.file.open('rb') as source:
			word = WordDocument(source)
	except FileNotFoundError:
		return HttpResponse('No se encontró el archivo del documento.', status=404)
	output = BytesIO()
	styles = getSampleStyleSheet()
	body_style = ParagraphStyle('DocumentBody', parent=styles['BodyText'], fontSize=9.5, leading=14, spaceAfter=7)
	story = []
	image_parts = {relationship.rId: relationship.target_part.blob for relationship in word.part.rels.values() if 'image' in relationship.reltype}
	chart_parts = {}
	for relationship in word.part.rels.values():
		if 'chart' in relationship.reltype:
			chart = render_docx_chart(relationship.target_part.blob)
			if chart:
				chart_output = BytesIO()
				chart.save(chart_output, format='PNG')
				chart_parts[relationship.rId] = chart_output.getvalue()
	def add_images(paragraph):
		for blip in paragraph._p.xpath('.//a:blip'):
			blob = image_parts.get(blip.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed'))
			if blob:
				with PillowImage.open(BytesIO(blob)) as image:
					width, height = image.size
				scale = min((175 * mm) / width, (220 * mm) / height, 1)
				story.extend([Spacer(1, 6), PdfImage(BytesIO(blob), width=width * scale, height=height * scale), Spacer(1, 8)])
		for chart_node in paragraph._p.xpath('.//c:chart'):
			blob = chart_parts.get(chart_node.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id'))
			if blob:
				story.extend([Spacer(1, 10), PdfImage(BytesIO(blob), width=180 * mm, height=97.5 * mm), Spacer(1, 16)])

	def add_table(table):
		rows = [[cell.text.strip().replace('\n', ' / ') for cell in row.cells] for row in table.rows]
		if rows:
			rendered = Table(rows, repeatRows=1, hAlign='LEFT')
			rendered.setStyle(TableStyle([('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#c9ed57')), ('GRID', (0, 0), (-1, -1), .45, colors.HexColor('#9da69d')), ('FONTSIZE', (0, 0), (-1, -1), 8.5), ('LEADING', (0, 0), (-1, -1), 11), ('TOPPADDING', (0, 0), (-1, -1), 6), ('BOTTOMPADDING', (0, 0), (-1, -1), 6), ('LEFTPADDING', (0, 0), (-1, -1), 5), ('RIGHTPADDING', (0, 0), (-1, -1), 5), ('VALIGN', (0, 0), (-1, -1), 'TOP')]))
			story.extend([Spacer(1, 8), rendered, Spacer(1, 12)])

	for element in word.element.body.iterchildren():
		if isinstance(element, CT_P):
			paragraph = WordParagraph(element, word)
			text = paragraph.text.strip()
			if text:
				story.append(Paragraph(html.escape(text), body_style))
			add_images(paragraph)
		elif isinstance(element, CT_Tbl):
			add_table(WordTable(element, word))
	SimpleDocTemplate(output, pagesize=A4, rightMargin=16 * mm, leftMargin=16 * mm, topMargin=15 * mm, bottomMargin=15 * mm).build(story)
	response = HttpResponse(output.getvalue(), content_type='application/pdf')
	disposition = 'attachment' if download else 'inline'
	response['Content-Disposition'] = f'{disposition}; filename="{Path(document.file.name).stem}.pdf"'
	return response


def build_preview(document):
	"""Extract Office content so the document can be read without downloading it."""
	extension = Path(document.file.name).suffix.lower() if document.file else '.pdf'
	preview = {'kind': extension.lstrip('.'), 'content': None, 'error': None}
	if not document.file:
		preview['error'] = 'El archivo no está disponible en la carpeta de documentos.'
		return preview
	try:
		if extension == '.pdf':
			preview['content'] = {'title': document.title, 'message': 'Este documento está disponible en PDF. Puede abrirse desde la vista de detalle.'}
		elif extension == '.docx':
			with document.file.open('rb') as source:
				word = WordDocument(source)
			images = []
			for relationship in word.part.rels.values():
				if 'image' in relationship.reltype:
					content_type = relationship.target_part.content_type
					encoded = base64.b64encode(relationship.target_part.blob).decode('ascii')
					images.append({'src': f'data:{content_type};base64,{encoded}'})
			preview['content'] = {
				'paragraphs': [paragraph.text for paragraph in word.paragraphs if paragraph.text.strip()],
				'tables': [[[cell.text for cell in row.cells] for row in table.rows] for table in word.tables],
				'images': images,
			}
		elif extension == '.xlsx':
			with document.file.open('rb') as source:
				workbook = load_workbook(source, data_only=True, read_only=True)
				try:
					dashboard = next((sheet for sheet in workbook.worksheets if sheet.title.strip().upper() == 'DASHBOARD'), None)
					if dashboard:
						preview['content'] = [{'name': dashboard.title, 'rows': [list(row) for row in dashboard.iter_rows(max_row=60, max_col=20, values_only=True)]}]
						preview['dashboard'] = dashboard_preview_data(document)
						preview['dashboard_only'] = True
					else:
						preview['content'] = []
				finally:
					workbook.close()
		elif extension == '.pptx':
			with document.file.open('rb') as source:
				presentation = Presentation(source)
			slides = []
			for index, slide in enumerate(presentation.slides, 1):
				titles = []
				texts = []
				for shape in slide.shapes:
					if not getattr(shape, 'has_text_frame', False):
						continue
					for paragraph in shape.text_frame.paragraphs:
						text = paragraph.text.strip()
						if not text:
							continue
						if len(titles) == 0 and not texts:
							titles.append(text)
						texts.append(text)
				if texts:
					unique_texts = []
					seen = set()
					for text in texts:
						if text and text not in seen:
							unique_texts.append(text)
							seen.add(text)
					slides.append({'number': index, 'title': titles[0] if titles else '', 'texts': unique_texts[:4]})
			preview['content'] = slides
	except FileNotFoundError:
		preview['error'] = 'El archivo no está disponible en el almacenamiento. Vuelve a cargarlo desde el administrador.'
	except (OSError, ValueError, KeyError, AttributeError):
		preview['error'] = 'No fue posible leer este archivo. Puedes abrirlo con su aplicación original.'

	return preview


def _extract_excel_dashboard(document):
	excel_source = document.file.open('rb')
	workbook = load_workbook(excel_source, data_only=True, read_only=True)
	keyed = {
		'CONSULTA 1': 'annual',
		'CONSULTA 2': 'monthly',
		'CONSULTA 3': 'weapons',
		'CONSULTA 4': 'ages',
		'CONSULTA 5': 'days',
		'CONSULTA 6': 'regions',
	}
	data = {'annual': [], 'monthly': [], 'days': [], 'weapons': [], 'ages': [], 'regions': [], 'regional_annual': [], 'gender': [], 'kpis': [], 'facets': []}

	for sheet_name, key in keyed.items():
		if sheet_name not in workbook.sheetnames:
			continue
		worksheet = workbook[sheet_name]
		rows = list(worksheet.iter_rows(values_only=True))
		if not rows:
			continue
		if key == 'ages':
			header = next((row for row in rows if row and row[0] == 'Etiquetas de fila'), None)
			total_row = next((row for row in reversed(rows) if row and str(row[0]).strip().upper() == 'TOTAL GENERAL'), None)
			if header:
				for index, label in enumerate(header[1:], start=1):
					if label is None or str(label).strip().upper() == 'TOTAL GENERAL':
						continue
					value = total_row[index] if total_row and len(total_row) > index else sum(
						_dashboard_number(row[index]) for row in rows if len(row) > index and isinstance(row[0], int)
					)
					data[key].append({'label': str(label).strip(), 'value': _dashboard_display(value)})
			continue
		if key == 'regions':
			header = next((row for row in rows if row and row[0] == 'Etiquetas de fila'), None)
			if not header:
				continue
			region_columns = [(index, str(label).strip()) for index, label in enumerate(header[1:], start=1) if label and str(label).strip().upper() != 'TOTAL GENERAL']
			for row in rows:
				if not row or not isinstance(row[0], int):
					continue
				for index, region in region_columns:
					if len(row) > index and row[index] is not None:
						data['regional_annual'].append({'year': str(row[0]), 'region': region, 'value': _dashboard_display(row[index])})
			total_row = next((row for row in reversed(rows) if row and str(row[0]).strip().upper() == 'TOTAL GENERAL'), None)
			if total_row:
				for index, region in region_columns:
					if len(total_row) > index and total_row[index] is not None:
						data['regions'].append({'label': region, 'value': _dashboard_display(total_row[index])})
			continue
		if key == 'annual':
			for row in rows[2:]:
				if not row or row[0] is None:
					continue
				label = row[0]
				value = row[1] if len(row) > 1 else None
				if isinstance(label, int) and value is not None:
					data[key].append({'label': str(label), 'value': _dashboard_display(value)})
			continue
			
		for row in rows:
			if not row or row[0] is None:
				continue
			if isinstance(row[0], str) and 'Etiquetas de fila' in row[0]:
				continue
			if isinstance(row[0], str) and row[0].upper().startswith('TOTAL'):
				continue
			if key == 'monthly':
				if row[0] == 'NOM MES':
					continue
				values = [cell for cell in row[1:7] if cell is not None]
				if values:
					series_total = sum(_dashboard_number(v) for v in values[:5])
					data[key].append({'label': str(row[0]), 'value': _dashboard_display(series_total)})
				continue
			if key == 'weapons':
				if len(row) < 5 or row[0] is None:
					continue
				label = str(row[0]).strip()
				if label.lower() == 'etiquetas de fila':
					continue
				if not label or label.upper().startswith('TOTAL'):
					continue
				value = row[-1] if row[-1] is not None else 0
				if value is not None:
					data[key].append({'label': label, 'value': _dashboard_display(value)})
				continue
			if key == 'ages':
				if len(row) < 6 or row[0] is None or row[0] == 'Etiquetas de fila':
					continue
				if row[0] == 'Total general':
					continue
				value = row[-1]
				if value is not None:
					data[key].append({'label': str(row[0]), 'value': _dashboard_display(value)})
				continue
			if key == 'days':
				if len(row) < 8 or row[0] is None:
					continue
				if row[0] == 'Etiquetas de fila':
					continue
				if row[0] == 'Total general':
					continue
				value = row[-1]
				if value is not None:
					data[key].append({'label': str(row[0]), 'value': _dashboard_display(value)})
				continue

	gender_sheet = workbook['CONSULTA 3'] if 'CONSULTA 3' in workbook.sheetnames else None
	if gender_sheet:
		rows = list(gender_sheet.iter_rows(values_only=True))
		gender_totals = {}
		for row in rows:
			if len(row) < 5:
				continue
			if row[0] is None or row[0] == 'Etiquetas de fila':
				continue
			if row[0].upper().startswith('TOTAL'):
				continue
			for idx, label in enumerate(['FEMENINO', 'MASCULINO', 'NO REPORTADO'], start=1):
				if len(row) <= idx:
					continue
				if row[idx] is not None:
					gender_totals[label] = gender_totals.get(label, 0) + _dashboard_number(row[idx])
		for label, value in gender_totals.items():
			data['gender'].append({'label': label, 'value': _dashboard_display(value)})

	detail_sheet = workbook['LESIONES PERSONALES UNIFICADA02'] if 'LESIONES PERSONALES UNIFICADA02' in workbook.sheetnames else None
	if detail_sheet:
		facets = defaultdict(float)
		for row in detail_sheet.iter_rows(min_row=2, values_only=True):
			if len(row) <= 22 or row[22] is None or row[2] is None or row[4] is None or row[9] is None or row[17] is None or row[19] is None:
				continue
			# col[2]=REGION, col[4]=DEPARTAMENTO, col[9]=AÑO, col[11]=MES(num), col[12]=NOM_MES,
			# col[17]=GENERO, col[19]=GRUPO_EDAD, col[0]=ARMAS_MEDIOS, col[22]=CANTIDAD
			month_num = str(row[11]).strip() if row[11] is not None else ''
			month_name = str(row[12]).strip() if row[12] is not None else ''
			key = (
				str(row[9]).strip(),          # year
				str(row[2]).strip(),          # region
				str(row[4]).strip(),          # dept
				month_num,                    # month number
				month_name,                   # month name
				str(row[17]).strip(),         # gender
				str(row[19]).strip(),         # age group
				str(row[0]).strip(),          # weapon
			)
			facets[key] += _dashboard_number(row[22])
		data['facets'] = [
			{
				'year': year, 'region': region, 'dept': dept,
				'month': month_num, 'month_name': month_name,
				'gender': gender, 'age': age, 'weapon': weapon,
				'value': _dashboard_display(value),
			}
			for (year, region, dept, month_num, month_name, gender, age, weapon), value in facets.items()
		]

	annual_sheet = workbook['CONSULTA 1']
	annual_rows = list(annual_sheet.iter_rows(values_only=True))
	annual_values = []
	for row in annual_rows[2:]:
		if len(row) > 1 and row[0] is not None and isinstance(row[0], int):
			annual_values.append((row[0], _dashboard_number(row[1])))
	if annual_values:
		top_year = max(annual_values, key=lambda item: item[1])[0]
		total_cases = sum(value for _, value in annual_values)
		data['kpis'] = [
			{'label': 'Total casos', 'value': _dashboard_display(total_cases)},
			{'label': 'Año con más casos', 'value': str(top_year)},
			{'label': 'Municipios', 'value': '1.020'},
			{'label': 'Hechos analizados', 'value': '332.390'},
		]
	workbook.close()
	excel_source.close()
	return _enrich_dashboard_data(data)


def _enrich_dashboard_data(data):
	if not data or not isinstance(data, dict):
		return data
	annual = data.get('annual', [])
	if annual:
		to_num = lambda v: _dashboard_number(v)
		total_victims = sum(to_num(item.get('value')) for item in annual)
		prev_val = None
		for item in annual:
			val = to_num(item.get('value'))
			share = (val / total_victims * 100) if total_victims else 0
			if 'detail' not in item or not item.get('detail'):
				item['detail'] = f'{share:.1f}%'.replace('.', ',')
			if prev_val is not None and prev_val > 0:
				growth = ((val - prev_val) / prev_val * 100)
				item['growth'] = f'{growth:+.1f}%'.replace('.', ',')
			else:
				item['growth'] = '—'
			prev_val = val
	return data


def dashboard_preview_data(document=None):
	if document is None:
		document = Document.objects.filter(document_type__in=['excel', 'xlsx'], is_published=True).first() or Document.objects.filter(document_type='word', is_published=True).first()
	if not document:
		return _enrich_dashboard_data({'annual': [], 'monthly': [], 'days': [], 'weapons': [], 'ages': [], 'regions': [], 'gender': [], 'kpis': [{'label': 'Total casos', 'value': '0'}, {'label': 'Año con más casos', 'value': '-'}, {'label': 'Municipios', 'value': '0'}]})
	if document.document_type in {'excel', 'xlsx'}:
		if not document.file:
			return _enrich_dashboard_data({'annual': [], 'monthly': [], 'days': [], 'weapons': [], 'ages': [], 'regions': [], 'gender': [], 'kpis': [{'label': 'Total casos', 'value': '0'}, {'label': 'Año con más casos', 'value': '-'}, {'label': 'Municipios', 'value': '0'}]})
		try:
			cache_key = (document.pk, document.file.name, document.file.size)
		except OSError:
			return _enrich_dashboard_data({'annual': [], 'monthly': [], 'days': [], 'weapons': [], 'ages': [], 'regions': [], 'gender': [], 'kpis': [{'label': 'Total casos', 'value': '0'}, {'label': 'Año con más casos', 'value': '-'}, {'label': 'Municipios', 'value': '0'}]})
		try:
			if cache_key not in _dashboard_cache or not _dashboard_cache[cache_key].get('facets'):
				_dashboard_cache.clear()
				precomputed_cache = settings.BASE_DIR / 'documentos' / 'dashboard_cache.json'
				if document.slug == 'lesiones-personales-en-colombia-20212025' and precomputed_cache.is_file():
					cached_data = json.loads(precomputed_cache.read_text(encoding='utf-8'))['data']
					if not isinstance(cached_data, dict) or not cached_data.get('facets'):
						raise ValueError(f'La caché de dashboard no tiene el formato esperado: {precomputed_cache}')
					_dashboard_cache[cache_key] = _enrich_dashboard_data(cached_data)
				else:
					_dashboard_cache[cache_key] = _extract_excel_dashboard(document)
			return _enrich_dashboard_data(_dashboard_cache[cache_key])
		except FileNotFoundError:
			return _enrich_dashboard_data({'annual': [], 'monthly': [], 'days': [], 'weapons': [], 'ages': [], 'regions': [], 'gender': [], 'kpis': [{'label': 'Total casos', 'value': '0'}, {'label': 'Año con más casos', 'value': '-'}, {'label': 'Municipios', 'value': '0'}]})
		except (OSError, ValueError, IndexError, PackageNotFoundError, BadZipFile, KeyError, TypeError, AttributeError):
			return _enrich_dashboard_data({'annual': [], 'monthly': [], 'days': [], 'weapons': [], 'ages': [], 'regions': [], 'gender': [], 'kpis': [{'label': 'Total casos', 'value': '0'}, {'label': 'Año con más casos', 'value': '-'}, {'label': 'Municipios', 'value': '0'}]})
	data = {'annual': [], 'monthly': [], 'days': [], 'weapons': [], 'ages': [], 'regions': [], 'gender': [], 'kpis': [{'label': 'Total casos', 'value': '495.272'}, {'label': 'Año con más casos', 'value': '2022'}, {'label': 'Municipios', 'value': '1.020'}, {'label': 'Hechos analizados', 'value': '332.390'}]}
	if not document.file or Path(document.file.name).suffix.lower() == '.pdf':
		return _enrich_dashboard_data(data)
	try:
		with document.file.open('rb') as source:
			word = WordDocument(source)
		for index, key in ((0, 'annual'), (1, 'monthly'), (2, 'weapons'), (3, 'ages'), (4, 'days'), (5, 'regions')):
			if len(word.tables) <= index:
				continue
			for row in word.tables[index].rows[1:]:
				values = [cell.text.strip().replace('\n', ' ') for cell in row.cells]
				if len(values) >= 2 and values[0]:
					data[key].append({'label': values[0], 'value': values[1], 'detail': values[2] if len(values) > 2 else ''})
		return _enrich_dashboard_data(data)
	except (OSError, ValueError, IndexError, PackageNotFoundError, BadZipFile, KeyError, TypeError, AttributeError):
		return _enrich_dashboard_data(data)


@cache_page(60 * 5)
def map_api(request):
	"""Endpoint JSON para el mapa coroplético de Colombia.

	Parámetros GET:
	  year   — Año (ej. 2023). Vacío = todos.
	  month  — Número de mes 1-12. Vacío = todos.
	  region — Nombre de región. Vacío = todas.
	  dept   — Nombre de departamento. Vacío = todos.
	  sex    — Género (MASCULINO / FEMENINO / NO REPORTADO). Vacío = todos.
	  age    — Grupo de edad. Vacío = todos.

	Respuesta JSON:
	  {
	    "dept_totals": {"ANTIOQUIA": 35000, ...},
	    "map_dept_totals": {"ANTIOQUIA": 35000, ...},
	    "total": 495272,
	    "map_total": 495272,
	    "filters": {years, months, regions, depts, genders, ages},
	    "active_filters": {year, month, region, dept, sex, age}
	  }
	"""
	statistics = dashboard_preview_data()
	facets = statistics.get('facets', [])

	# Leer filtros activos (normalizar a mayúsculas)
	f_year   = request.GET.get('year', '').strip()
	f_month  = request.GET.get('month', '').strip()
	f_region = request.GET.get('region', '').strip().upper()
	f_dept   = request.GET.get('dept', '').strip().upper()
	f_sex    = request.GET.get('sex', '').strip().upper()
	f_age    = request.GET.get('age', '').strip().upper()

	# Recorrer facets y acumular totales por departamento
	dept_totals = defaultdict(float)
	map_dept_totals = defaultdict(float)
	all_years, all_months, all_regions, all_depts, all_genders, all_ages = set(), set(), set(), set(), set(), set()

	for facet in facets:
		year    = str(facet.get('year', '')).strip()
		region  = str(facet.get('region', '')).strip().upper()
		dept    = str(facet.get('dept', '')).strip().upper()
		month   = str(facet.get('month', '')).strip()
		month_n = str(facet.get('month_name', '')).strip().upper()
		gender  = str(facet.get('gender', '')).strip().upper()
		age     = str(facet.get('age', '')).strip().upper()
		value   = _dashboard_number(facet.get('value', 0))

		# Recolectar valores únicos para los selectores
		if year:   all_years.add(year)
		if month_n: all_months.add((month, month_n))
		if region: all_regions.add(region)
		if dept:   all_depts.add(dept)
		if gender: all_genders.add(gender)
		if age:    all_ages.add(age)

		# Aplicar filtros
		if f_year   and year   != f_year:                    continue
		if f_month  and month  != f_month:                   continue
		if f_region and region != f_region:                  continue
		if f_sex    and gender != f_sex:                     continue
		if f_age    and age    != f_age:                     continue

		if dept:
			map_dept_totals[dept] += value
			if not f_dept or dept == f_dept:
				dept_totals[dept] += value

	total = sum(dept_totals.values())
	map_total = sum(map_dept_totals.values())

	# Ordenar meses por número
	months_sorted = sorted(all_months, key=lambda item: int(item[0]) if item[0].isdigit() else 99)

	return JsonResponse({
		'dept_totals': {k: int(v) for k, v in sorted(dept_totals.items(), key=lambda item: -item[1])},
		'map_dept_totals': {k: int(v) for k, v in sorted(map_dept_totals.items(), key=lambda item: -item[1])},
		'total': int(total),
		'map_total': int(map_total),
		'active_filters': {
			'year': f_year, 'month': f_month, 'region': f_region,
			'dept': f_dept, 'sex': f_sex, 'age': f_age,
		},
		'filters': {
			'years':   sorted(all_years, reverse=True),
			'months':  [{'num': m[0], 'name': m[1].title()} for m in months_sorted],
			'regions': sorted(all_regions),
			'depts':   sorted(all_depts),
			'genders': sorted(all_genders),
			'ages':    sorted(all_ages),
		},
	})

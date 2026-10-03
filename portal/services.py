from __future__ import annotations

import hashlib
import json
import logging
import re
import unicodedata
from zipfile import BadZipFile

from django.core.cache import cache
from docx import Document as WordDocument
from docx.opc.exceptions import PackageNotFoundError
from pypdf import PdfReader
from pypdf.errors import PdfReadError
from pptx import Presentation

from .utils import dashboard_number as _dashboard_number, dashboard_display as _dashboard_display

logger = logging.getLogger(__name__)
_SEARCH_STOPWORDS = {
    'al', 'como', 'con', 'cual', 'cuales', 'del', 'en', 'esta', 'este', 'fue',
    'hay', 'las', 'los', 'mas', 'para', 'por', 'que', 'sobre', 'una', 'uno',
    'entre', 'son',
}


def _search_normalize(text):
    normalized = unicodedata.normalize('NFKD', str(text).casefold())
    return ''.join(character for character in normalized if not unicodedata.combining(character))


def _search_tokens(text):
    return {
        token for token in re.findall(r'[a-z0-9]{2,}', _search_normalize(text))
        if token not in _SEARCH_STOPWORDS
    }


def _append_search_chunk(index, title, text, reference='', slug=''):
    text = ' '.join(str(text).split())
    if len(text) < 20:
        return
    text = text[:1800]
    tokens = _search_tokens(text)
    if not tokens:
        return
    index.append({
        'title': title,
        'text': text[:1800],
        'reference': reference,
        'slug': slug,
        'tokens': tokens,
    })


def _document_search_chunks(document, index):
    _append_search_chunk(index, document.title, document.summary, 'Resumen', document.slug)
    try:
        extension = document.file.name.rsplit('.', 1)[-1].lower() if document.file else ''
        if extension == 'docx':
            with document.file.open('rb') as source:
                word = WordDocument(source)
            for paragraph_number, paragraph in enumerate(word.paragraphs, 1):
                _append_search_chunk(index, document.title, paragraph.text, f'Párrafo {paragraph_number}', document.slug)
            for table_number, table in enumerate(word.tables, 1):
                for row_number, row in enumerate(table.rows, 1):
                    _append_search_chunk(
                        index, document.title,
                        ' | '.join(cell.text for cell in row.cells),
                        f'Tabla {table_number}, fila {row_number}', document.slug,
                    )
        elif extension == 'pptx':
            with document.file.open('rb') as source:
                presentation = Presentation(source)
            for slide_number, slide in enumerate(presentation.slides, 1):
                texts = []
                for shape in slide.shapes:
                    if getattr(shape, 'has_text_frame', False):
                        texts.append(shape.text)
                    if getattr(shape, 'has_table', False):
                        texts.extend(
                            ' | '.join(cell.text for cell in row.cells)
                            for row in shape.table.rows
                        )
                _append_search_chunk(index, document.title, ' '.join(texts), f'Diapositiva {slide_number}', document.slug)
        elif extension == 'pdf':
            with document.file.open('rb') as source:
                pdf = PdfReader(source)
                for page_number in range(min(len(pdf.pages), 100)):
                    _append_search_chunk(
                        index, document.title, pdf.pages[page_number].extract_text() or '',
                        f'Página {page_number + 1}', document.slug,
                    )
    except (OSError, ValueError, KeyError, AttributeError, BadZipFile, PackageNotFoundError, PdfReadError) as error:
        logger.warning('No se pudo indexar el documento %s para el chat: %s', document.title, error)


def _chat_index_signature(documents, statistics, bank):
    document_signature = []
    for document in documents:
        try:
            file_size = document.file.size if document.file else 0
        except OSError:
            file_size = None
        document_signature.append((
            document.pk, document.title, document.summary,
            document.file.name if document.file else '', file_size, document.slug,
        ))
    stat_signature = {
        key: [(item.get('label'), item.get('value')) for item in statistics.get(key, [])]
        for key in ('annual', 'monthly', 'days', 'weapons', 'ages', 'regions', 'gender')
    }
    bank_signature = [
        (category, question['id'], question['label'], question['answer'])
        for category, data in bank.items()
        for question in data['questions']
    ]
    serialized = json.dumps([document_signature, stat_signature, bank_signature], ensure_ascii=False, sort_keys=True)
    return hashlib.sha256(serialized.encode('utf-8')).hexdigest()


def _build_chat_search_index(documents, statistics, bank):
    index = []
    for document in documents:
        _document_search_chunks(document, index)

    for category in bank.values():
        for question in category['questions']:
            _append_search_chunk(
                index, 'Respuestas del asistente del proyecto',
                f"{question['label']} {question['answer']}", 'Banco de preguntas',
            )

    for key in ('annual', 'monthly', 'days', 'weapons', 'ages', 'regions', 'gender'):
        for item in statistics.get(key, []):
            _append_search_chunk(
                index, 'Dashboard de datos',
                f"{key}: {item.get('label', '')} — {item.get('value', '')} {item.get('detail', '')}",
                'Indicadores del dashboard',
            )
    return index


def search_chat_knowledge(query, documents, statistics, bank):
    signature = _chat_index_signature(documents, statistics, bank)
    cache_key = f'chat-search-index:{signature}'
    index = cache.get(cache_key)
    if index is None:
        index = _build_chat_search_index(documents, statistics, bank)
        cache.set(cache_key, index, timeout=60 * 60)

    query_tokens = _search_tokens(query)
    if not query_tokens:
        return []

    minimum_matches = min(2, len(query_tokens))
    normalized_query = _search_normalize(query)
    matches = []
    for entry in index:
        matched_tokens = query_tokens & entry['tokens']
        if len(matched_tokens) < minimum_matches:
            continue
        coverage = len(matched_tokens) / len(query_tokens)
        phrase_bonus = 0.15 if normalized_query in _search_normalize(entry['text']) else 0
        matches.append((coverage + phrase_bonus, entry))

    matches.sort(key=lambda match: match[0], reverse=True)
    results = []
    seen = set()
    for score, entry in matches:
        if score < 0.34:
            break
        identity = (entry['title'], entry['text'])
        if identity in seen:
            continue
        seen.add(identity)
        results.append({key: value for key, value in entry.items() if key != 'tokens'})
        if len(results) == 3:
            break
    return results


def get_question_bank():
    from .views import dashboard_preview_data

    statistics = dashboard_preview_data()
    annual = statistics.get('annual', [])
    regions = statistics.get('regions', [])
    gender = statistics.get('gender', [])
    ages = statistics.get('ages', [])
    peak = max(annual, key=lambda item: _dashboard_number(item.get('value')), default={'label': '2022', 'value': '110.410'})
    top_region = max(regions, key=lambda item: _dashboard_number(item.get('value')), default={'label': 'REGION ANDINA', 'value': '315.074'})
    top_age = max(ages, key=lambda item: _dashboard_number(item.get('value')), default={'label': 'ADULTOS', 'value': '455.211'})
    top_gender = max(gender, key=lambda item: _dashboard_number(item.get('value')), default={'label': 'MASCULINO', 'value': '298.881'})

    return {
        'dashboard': {
            'label': 'Cifras y Dashboard',
            'description': 'Cifras oficiales, víctimas vs. casos y hallazgos visuales.',
            'questions': [
                {'id': 'dashboard-overview', 'label': '¿Qué analiza el dashboard?', 'answer': 'El dashboard organiza la información de lesiones personales en tres dimensiones: tiempo (2021-2025), población (género, edad y armas) y territorio (regiones y municipios). Permite cruzar variables y comprender el comportamiento del fenómeno en Colombia.'},
                {'id': 'dashboard-cases-vs-victims', 'label': '¿Cuál es la diferencia entre casos y víctimas?', 'answer': 'El observatorio distingue dos métricas clave: los casos o hechos delictivos (332.390 eventos registrados) y las víctimas (495.272 personas lesionadas en las tablas operativas). En promedio, cada caso registrado involucra a 1,49 víctimas.'},
                {'id': 'dashboard-peak', 'label': '¿Qué año presenta el valor más alto y cuál es la tendencia?', 'answer': f"El año con mayor afectación fue {peak.get('label')}, con {peak.get('value')} víctimas registradas. A partir de 2022 se observa una disminución sostenida hasta 2025 (89.714 víctimas), equivalente a una reducción del -18,7% respecto al pico histórico."},
                {'id': 'dashboard-region', 'label': '¿Qué región concentra más registros?', 'answer': f"La Región Andina concentra el mayor volumen con {top_region.get('value')} registros agregados (63,6% del total nacional), seguida por la Región Caribe (15,3%) y la Región Pacífica (14,8%)."},
                {'id': 'dashboard-weapons', 'label': '¿Cuáles son los medios y armas más empleados?', 'answer': 'Los elementos contundentes (44,4% / 220.082 víctimas) y las agresiones sin empleo de armas (39,0% / 193.159 víctimas) representan el 83,4% de los hechos, evidenciando el predominio de riñas interpersonales. Las armas blancas alcanzan el 10,6% y las de fuego el 4,2%.'},
                {'id': 'dashboard-days', 'label': '¿Qué días de la semana son los más críticos?', 'answer': 'El domingo es el día con mayor afectación, concentrando 119.048 víctimas (24,0%), seguido del sábado con 81.434 (16,4%). El fin de semana acumula el 40,4% de todas las lesiones registradas.'},
            ],
        },
        'documentation': {
            'label': 'Documentación y Metodología',
            'description': 'Metodología, fuentes institucionales y fuentes de datos.',
            'questions': [
                {'id': 'documentation-purpose', 'label': '¿Cuál es el propósito del proyecto?', 'answer': 'Identificar, cuantificar y visualizar patrones temporales, demográficos y territoriales de lesiones personales en Colombia entre 2021 y 2025 para aportar evidencia a la política pública de seguridad y convivencia ciudadana.'},
                {'id': 'documentation-method', 'label': '¿Qué metodología utiliza?', 'answer': 'La documentación integra, depura y explora microdatos oficiales para construir indicadores comparables. Aplica análisis temporal, poblacional y territorial con trazabilidad metodológica completa.'},
                {'id': 'documentation-results', 'label': '¿Cuáles son las cifras principales del informe?', 'answer': 'El informe documenta 332.390 casos delictivos consolidados y 495.272 víctimas acumuladas entre 2021 y 2025. El dashboard interactivo permite explorar la serie anual y los cortes temáticos.'},
                {'id': 'documentation-source', 'label': '¿Cuál es la fuente institucional?', 'answer': 'La fuente institucional oficial es el Sistema de Información Estadístico, Delincuencial, Contravencional y Operativo (SIEDCO) de la Policía Nacional de Colombia, con cobertura en más de 1.020 municipios.'},
                {'id': 'documentation-resource', 'label': '¿Qué documentos están publicados?', 'answer': 'Están publicados tres recursos principales: el documento analítico en Word y PDF de lectura, el libro Excel con los datos operativos del dashboard, y la presentación general del proyecto.'},
            ],
        },
        'presentation': {
            'label': 'Presentación y Conclusiones',
            'description': 'Resumen ejecutivo, perfiles vulnerables y recomendaciones.',
            'questions': [
                {'id': 'presentation-findings', 'label': '¿Cuáles son los hallazgos para destacar?', 'answer': f"Destaca el pico anual en {peak.get('label')}, el descenso posterior de -18,7%, la concentración del 63,6% en la Región Andina, la predominancia de adultos (91,9%) y el hecho de que más del 83% de las agresiones provienen de riñas y golpes contundentes sin armas de fuego."},
                {'id': 'presentation-profile', 'label': '¿Cuál es el perfil demográfico predominante?', 'answer': f"El grupo más representado corresponde a {top_age.get('label', '').title()} con {top_age.get('value')} víctimas (91,9%). A nivel de género, los hombres representan el 60,3% ({top_gender.get('value')} víctimas) y las mujeres el 39,6% (196.115 víctimas)."},
                {'id': 'presentation-explain', 'label': '¿Cómo puedo explicar el proyecto?', 'answer': 'Puedes presentarlo como un observatorio que transforma microdatos policiales en una lectura clara de tendencias, perfiles y territorios para orientar conversaciones de prevención, seguridad y justicia.'},
                {'id': 'presentation-audience', 'label': '¿Para quién es útil este observatorio?', 'answer': 'Para investigadores, equipos de formulación de política pública, autoridades policiales, periodistas y analistas que requieran evidencia trazable sobre violencia interpersonal en Colombia.'},
                {'id': 'presentation-use', 'label': '¿Cómo se usa el dashboard en una presentación?', 'answer': 'Inicia con la tendencia anual, aplica filtros por región o grupo de edad para evidenciar patrones específicos y contrasta la lectura con la metodología y la fuente institucional.'},
            ],
        },
    }


def get_chat_categories():
    bank = get_question_bank()
    return [{
        'id': key,
        'label': value['label'],
        'description': value['description'],
        'questions': [{'id': item['id'], 'label': item['label']} for item in value['questions']],
    } for key, value in bank.items()]


def get_quick_questions():
    return [
        {'id': 'dashboard-cases-vs-victims', 'category': 'dashboard', 'label': '¿Diferencia entre casos y víctimas?'},
        {'id': 'dashboard-overview', 'category': 'dashboard', 'label': '¿Qué analiza el dashboard?'},
        {'id': 'dashboard-peak', 'category': 'dashboard', 'label': '¿Año con más víctimas y tendencia?'},
        {'id': 'presentation-findings', 'category': 'presentation', 'label': '¿Cuáles son los hallazgos para destacar?'},
    ]


def build_homepage_summary(statistics):
    annual = statistics.get('annual', [])
    regions = statistics.get('regions', [])
    gender = statistics.get('gender', [])
    ages = statistics.get('ages', [])
    to_number = lambda value: float(str(value or '0').replace('.', '').replace(',', '.'))
    annual_values = [(item.get('label', ''), to_number(item.get('value'))) for item in annual]
    peak_year, peak_value = max(annual_values, key=lambda item: item[1], default=('-', 0))
    first_value = annual_values[0][1] if annual_values else 0
    last_value = annual_values[-1][1] if annual_values else 0
    change = ((last_value - first_value) / first_value * 100) if first_value else 0
    peak_reduction = ((last_value - peak_value) / peak_value * 100) if peak_value else 0

    total_victims = sum(val for _, val in annual_values)
    top_region_item = max(regions, key=lambda item: to_number(item.get('value')), default={'label': 'Sin datos', 'value': '0'})
    top_region_num = to_number(top_region_item.get('value'))
    top_region_share = f'{(top_region_num / total_victims * 100):.1f}%'.replace('.', ',') if total_victims else '0%'

    top_gender_item = max(gender, key=lambda item: to_number(item.get('value')), default={'label': 'Sin datos', 'value': '0'})
    top_gender_num = to_number(top_gender_item.get('value'))
    top_gender_share = f'{(top_gender_num / total_victims * 100):.1f}%'.replace('.', ',') if total_victims else '0%'

    top_age_item = max(ages, key=lambda item: to_number(item.get('value')), default={'label': 'Sin datos', 'value': '0'})
    top_age_num = to_number(top_age_item.get('value'))
    top_age_share = f'{(top_age_num / total_victims * 100):.1f}%'.replace('.', ',') if total_victims else '0%'

    chart_points = []
    if len(annual_values) >= 2:
        x_coords = [60, 160, 260, 360, 460]
        y_min, y_max = 75, 190
        max_val = max(v for _, v in annual_values)
        min_val = min(v for _, v in annual_values)
        span = (max_val - min_val) or 1
        for i, (yr, val) in enumerate(annual_values):
            x = x_coords[i] if i < len(x_coords) else int(round(60 + i * 100))
            y = int(round(y_max - ((val - min_val) / span) * (y_max - y_min)))
            label_y = int(y - 14)
            chart_points.append({
                'cx': int(x),
                'cy': int(y),
                'label_y': int(label_y),
                'year': yr,
                'value': _dashboard_display(val),
            })
    polyline_str = ' '.join(f"{p['cx']},{p['cy']}" for p in chart_points)

    return {
        'peak_year': peak_year,
        'peak_value': _dashboard_display(peak_value),
        'change': f'{change:+.1f}%'.replace('.', ','),
        'peak_reduction': f'{peak_reduction:+.1f}%'.replace('.', ','),
        'top_region': {
            'label': top_region_item.get('label', 'Sin datos'),
            'value': top_region_item.get('value', '0'),
            'share': top_region_share,
        },
        'top_gender': {
            'label': top_gender_item.get('label', 'Sin datos'),
            'value': top_gender_item.get('value', '0'),
            'share': top_gender_share,
        },
        'top_age': {
            'label': top_age_item.get('label', 'Sin datos'),
            'value': top_age_item.get('value', '0'),
            'share': top_age_share,
        },
        'region_count': len(regions),
        'total_victims': _dashboard_display(total_victims) if total_victims else '495.272',
        'total_cases': '332.390',
        'chart_points': chart_points,
        'chart_polyline': polyline_str,
    }


def get_answer_for_question(category_id: str, question_id: str):
    bank = get_question_bank()
    category_key = str(category_id or '')
    category = bank.get(category_key)
    if not category:
        return None
    question = next((item for item in category['questions'] if item['id'] == str(question_id)), None)
    return question['answer'] if question else None

from io import BytesIO
from pathlib import Path
import shutil
import tempfile
from unittest.mock import patch

from django.conf import settings
from django.core.cache import cache
from django.core import management
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.files.storage import default_storage
from django.test import TestCase, override_settings
from django.urls import reverse
from docx import Document as WordDocument
from openpyxl import Workbook

from .models import Document
from .views import dashboard_preview_data


@override_settings(SECURE_SSL_REDIRECT=False)
class DocumentViewsTests(TestCase):
	@classmethod
	def setUpClass(cls):
		super().setUpClass()
		cls.media_directory = tempfile.mkdtemp(prefix='portal-tests-')
		cls.media_override = override_settings(MEDIA_ROOT=cls.media_directory)
		cls.media_override.enable()

	@classmethod
	def tearDownClass(cls):
		cls.media_override.disable()
		shutil.rmtree(cls.media_directory, ignore_errors=True)
		super().tearDownClass()

	def tearDown(self):
		cache.clear()
		super().tearDown()

	def _create_word_file(self, name='sample.docx'):
		doc = WordDocument()
		doc.add_paragraph('Documento de prueba')
		buffer = BytesIO()
		doc.save(buffer)
		buffer.seek(0)
		return SimpleUploadedFile(name, buffer.getvalue(), content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document')

	def test_home_page_loads(self):
		response = self.client.get(reverse('portal:home'))
		self.assertEqual(response.status_code, 200)
		self.assertNotContains(response, 'id="colombia-map-container"')

	def test_home_shows_documents_published_after_initial_visit(self):
		self.client.get(reverse('portal:home'))
		Document.objects.create(
			title='Recurso recién publicado',
			slug='recurso-recien-publicado',
			document_type='excel',
			summary='Resumen del recurso',
			file=SimpleUploadedFile('recurso.xlsx', b'contenido'),
			is_published=True,
		)

		response = self.client.get(reverse('portal:home'))

		self.assertContains(response, 'Recurso recién publicado')

	def test_published_document_file_is_served(self):
		document = Document.objects.create(
			title='Archivo público',
			slug='archivo-publico',
			document_type='excel',
			summary='Resumen',
			file=SimpleUploadedFile('archivo.xlsx', b'contenido-del-excel'),
			is_published=True,
		)

		response = self.client.get(document.file.url)

		self.assertEqual(response.status_code, 200)
		self.assertEqual(b''.join(response.streaming_content), b'contenido-del-excel')

	def test_document_pdf_can_be_embedded_on_same_site(self):
		document = Document.objects.create(
			title='Documento con PDF',
			slug='documento-con-pdf',
			document_type='word',
			summary='Resumen',
			file=self._create_word_file('documento.docx'),
			pdf_file=SimpleUploadedFile('documento.pdf', b'%PDF contenido', content_type='application/pdf'),
			is_published=True,
		)

		response = self.client.get(reverse('portal:document_pdf', args=[document.slug]))

		self.assertEqual(response.status_code, 200)
		self.assertEqual(response['Content-Type'], 'application/pdf')
		self.assertEqual(response['X-Frame-Options'], 'SAMEORIGIN')

	def test_published_document_media_is_downloadable(self):
		document = Document.objects.create(
			title='Archivo público',
			slug='archivo-publico',
			document_type='excel',
			summary='Resumen',
			file=SimpleUploadedFile('archivo.xlsx', b'contenido-del-excel'),
			is_published=True,
		)

		response = self.client.get(document.file.url)

		self.assertEqual(response.status_code, 200)
		self.assertEqual(b''.join(response.streaming_content), b'contenido-del-excel')

	def test_map_page_loads(self):
		response = self.client.get(reverse('portal:map'))
		self.assertRedirects(response, reverse('portal:territory'))

	def test_territory_landing_page_loads(self):
		response = self.client.get(reverse('portal:territory'))
		self.assertEqual(response.status_code, 200)
		self.assertContains(response, 'Mapa interactivo')

	def test_findings_page_loads(self):
		response = self.client.get(reverse('portal:findings'))
		self.assertEqual(response.status_code, 200)
		self.assertContains(response, 'Hallazgos principales')

	def test_findings_page_keeps_resource_download_area(self):
		Document.objects.create(
			title='Recurso descargable',
			slug='recurso-descargable',
			document_type='word',
			summary='Resumen del recurso',
			file=self._create_word_file('recurso-descargable.docx'),
			is_published=True,
		)
		response = self.client.get(reverse('portal:findings'))
		self.assertContains(response, 'Descargar')
		self.assertContains(response, 'Ver documento')

	def test_findings_shows_documents_published_after_initial_visit(self):
		self.client.get(reverse('portal:findings'))
		Document.objects.create(
			title='Hallazgo recién publicado',
			slug='hallazgo-recien-publicado',
			document_type='excel',
			summary='Resumen del recurso',
			file=SimpleUploadedFile('hallazgo.xlsx', b'contenido'),
			is_published=True,
		)

		response = self.client.get(reverse('portal:findings'))

		self.assertContains(response, 'Hallazgo recién publicado')

	def test_methodology_page_uses_glossary_content(self):
		response = self.client.get(reverse('portal:methodology'))
		self.assertEqual(response.status_code, 200)
		self.assertContains(response, 'Glosario del informe')
		self.assertContains(response, 'Código DANE')

	def test_dashboard_exposes_filter_tools(self):
		response = self.client.get(reverse('portal:dashboard'))
		self.assertEqual(response.status_code, 200)
		self.assertContains(response, 'Aplicación práctica')
		self.assertContains(response, 'Generar informe')
		self.assertContains(response, 'Descargar PDF')
		self.assertContains(response, 'Descargar TXT')
		self.assertContains(response, 'Exportar dashboard PDF')
		self.assertContains(response, 'Modo presentación')
		self.assertContains(response, 'informe-generado')
		self.assertContains(response, 'Restablecer')
		self.assertContains(response, 'Exportar CSV')
		self.assertContains(response, 'Copiar enlace')
		self.assertContains(response, 'Informes que ya')

	def test_dashboard_api_returns_data(self):
		response = self.client.get(reverse('portal:dashboard_api'))
		self.assertEqual(response.status_code, 200)
		self.assertIn('annual', response.json())

	@patch('portal.views._extract_excel_dashboard', side_effect=AssertionError('No debe volver a analizarse el Excel publicado en cada arranque.'))
	def test_repository_excel_uses_precomputed_dashboard_cache(self, extract_excel):
		document = Document.objects.create(
			title='Excel versionado',
			slug='lesiones-personales-en-colombia-20212025',
			document_type='excel',
			summary='Base de datos',
			file=SimpleUploadedFile('archivo.xlsx', b'excel'),
			is_published=True,
		)

		data = dashboard_preview_data(document)

		self.assertEqual(data['annual'][0]['label'], '2021')
		self.assertTrue(data['facets'])
		extract_excel.assert_not_called()

	def test_generated_report_pdf_returns_pdf(self):
		response = self.client.post(reverse('portal:generated_report_pdf'), data={'report_text': 'INFORME TÉCNICO\n\nVíctimas: 495.272'}, content_type='application/json')
		self.assertEqual(response.status_code, 200)
		self.assertEqual(response['Content-Type'], 'application/pdf')
		self.assertIn('informe-dashboard.pdf', response['Content-Disposition'])

	def test_report_page_loads(self):
		response = self.client.get(reverse('portal:report'))
		self.assertEqual(response.status_code, 200)
		self.assertContains(response, 'Informe ejecutivo')
		self.assertContains(response, 'Medio principal')

	def test_methodology_exposes_dynamic_glossary_tools(self):
		response = self.client.get(reverse('portal:methodology'))
		self.assertContains(response, 'data-glossary-search')
		self.assertContains(response, '36 categorías de armas y medios')

	def test_production_security_settings_are_configurable(self):
		self.assertTrue(hasattr(settings, 'ADMIN_URL'))
		self.assertTrue(hasattr(settings, 'CSRF_TRUSTED_ORIGINS'))
		self.assertIsInstance(settings.CSRF_TRUSTED_ORIGINS, list)
		self.assertIsInstance(settings.DATA_UPLOAD_MAX_MEMORY_SIZE, int)
		self.assertIsInstance(settings.FILE_UPLOAD_MAX_MEMORY_SIZE, int)
		self.assertLess(settings.FILE_UPLOAD_MAX_MEMORY_SIZE, settings.MAX_UPLOAD_SIZE)

	def test_project_qr_returns_png(self):
		response = self.client.get(reverse('portal:project_qr'))
		self.assertEqual(response.status_code, 200)
		self.assertEqual(response['Content-Type'], 'image/png')
		self.assertIn('no-store', response['Cache-Control'])

	def test_project_qr_uses_configured_public_url(self):
		with override_settings(PROJECT_PUBLIC_URL='http://192.168.1.25:8000'):
			response = self.client.get(reverse('portal:home'))
		self.assertContains(response, 'http://192.168.1.25:8000/')

	def test_map_api_returns_territorial_payload(self):
		response = self.client.get(reverse('portal:map_api'))
		payload = response.json()
		self.assertEqual(response.status_code, 200)
		self.assertIn('dept_totals', payload)
		self.assertIn('map_dept_totals', payload)
		self.assertIn('filters', payload)
		self.assertIn('active_filters', payload)

	@patch('portal.views.dashboard_preview_data')
	def test_map_api_keeps_other_department_counts_available_for_hover(self, dashboard_data):
		dashboard_data.return_value = {
			'facets': [
				{'year': '2025', 'region': 'REGION ANDINA', 'dept': 'BOYACÁ', 'month': '1', 'month_name': 'ENERO', 'gender': 'MASCULINO', 'age': 'ADULTOS', 'value': '20.580'},
				{'year': '2025', 'region': 'REGION ANDINA', 'dept': 'CUNDINAMARCA', 'month': '1', 'month_name': 'ENERO', 'gender': 'MASCULINO', 'age': 'ADULTOS', 'value': '150.124'},
				{'year': '2024', 'region': 'REGION ANDINA', 'dept': 'BOYACÁ', 'month': '1', 'month_name': 'ENERO', 'gender': 'MASCULINO', 'age': 'ADULTOS', 'value': '10'},
			],
		}

		response = self.client.get(reverse('portal:map_api'), {'dept': 'BOYACÁ', 'year': '2025'})
		payload = response.json()

		self.assertEqual(payload['dept_totals'], {'BOYACÁ': 20580})
		self.assertEqual(payload['total'], 20580)
		self.assertEqual(payload['map_dept_totals'], {'CUNDINAMARCA': 150124, 'BOYACÁ': 20580})
		self.assertEqual(payload['map_total'], 170704)
		self.assertEqual(payload['active_filters']['dept'], 'BOYACÁ')

	def test_chat_lists_question_categories(self):
		response = self.client.get(reverse('portal:chat'))
		self.assertEqual(response.status_code, 200)
		self.assertEqual({category['id'] for category in response.json()['categories']}, {'dashboard', 'documentation', 'presentation'})
		self.assertIn('quick_questions', response.json())
		self.assertTrue(response.json()['quick_questions'])

	def test_chat_answers_selected_question_locally(self):
		response = self.client.post(reverse('portal:chat'), data={'category': 'dashboard', 'question': 'dashboard-overview'}, content_type='application/json')
		self.assertEqual(response.status_code, 200)
		self.assertIn('tres dimensiones', response.json()['answer'])

	def test_chat_rejects_unknown_question(self):
		response = self.client.post(reverse('portal:chat'), data={'category': 'dashboard', 'question': 'unknown'}, content_type='application/json')
		self.assertEqual(response.status_code, 400)

	def test_unpublished_document_is_hidden_from_home(self):
		Document.objects.create(
			title='Documento privado',
			slug='documento-privado',
			document_type='word',
			summary='Resumen privado',
			file=self._create_word_file('documento-privado.docx'),
			is_published=False,
		)
		response = self.client.get(reverse('portal:home'))
		self.assertNotContains(response, 'Documento privado')

	def test_published_document_is_visible_on_home(self):
		Document.objects.create(
			title='Documento público',
			slug='documento-publico',
			document_type='word',
			summary='Resumen público',
			file=self._create_word_file('documento-publico.docx'),
			is_published=True,
		)
		response = self.client.get(reverse('portal:home'))
		self.assertContains(response, 'Documento público')

	def test_dashboard_reads_xlsx_values_instead_of_word_defaults(self):
		workbook = Workbook()
		annual = workbook.active
		annual.title = 'CONSULTA 1'
		annual.append(['SEGUIMIENTO ANUAL DE CASOS', None])
		annual.append([None, None])
		annual.append(['Etiquetas de fila', 'Suma de CANTIDAD '])
		annual.append([2021, 105316])
		annual.append([2022, 110410])
		annual.append([2023, 98886])
		annual.append([2024, 90946])
		annual.append([2025, 89714])
		annual.append(['Total general', 495272])

		gender = workbook.create_sheet('CONSULTA 3')
		gender.append(['ARMAS/MEDIO EMPLEADOS VS GENERO', None, None, None, None])
		gender.append([None, None, None, None, None])
		gender.append(['Suma de CANTIDAD ', 'Etiquetas de columna', None, None, None])
		gender.append(['Etiquetas de fila', 'FEMENINO', 'MASCULINO', 'NO REPORTADO ', 'Total general'])
		gender.append(['ARMA BLANCA', 11269, 41098, 13, 52380])
		gender.append(['ARMA DE FUEGO', 3085, 17809, 2, 20896])
		gender.append(['ARMA TRAUMATICA', 358, 2151, 2, 2511])

		ages = workbook.create_sheet('CONSULTA 4')
		ages.append(['DISTRIBUCION POR GRUPO DE EDAD', None, None, None, None, None])
		ages.append([None, None, None, None, None, None])
		ages.append(['Suma de CANTIDAD ', 'Etiquetas de columna', None, None, None, None])
		ages.append(['Etiquetas de fila', 'ADOLESCENTES', 'ADULTOS', 'MENORES', 'NO REPORTADO ', 'Total general'])
		ages.append([2021, 5273, 98179, 1786, 78, 105316])
		ages.append([2022, 8093, 100071, 2131, 115, 110410])
		ages.append([2023, 7169, 89933, 1612, 172, 98886])
		ages.append([2024, 6660, 82695, 1459, 132, 90946])
		ages.append([2025, 4447, 84333, 875, 59, 89714])
		ages.append(['Total general', 31642, 455211, 7863, 556, 495272])

		regions = workbook.create_sheet('CONSULTA 6')
		regions.append(['EVOLUCION DE CASOS POR REGION', None, None, None, None, None, None, None])
		regions.append([None, None, None, None, None, None, None, None])
		regions.append(['Suma de CANTIDAD ', 'Etiquetas de columna', None, None, None, None, None, None])
		regions.append(['Etiquetas de fila', 'REGION AMAZONIA', 'REGION ANDINA', 'REGION CARIBE', 'REGION INSULAR', 'REGION ORINOQUIA', 'REGION PACIFICA', 'Total general'])
		regions.append([2022, 1841, 68957, 18714, 343, 4851, 15704, 110410])
		regions.append([2021, 1606, 68503, 16220, 337, 4518, 14132, 105316])
		regions.append([2023, 1826, 62409, 15083, 221, 4468, 14879, 98886])

		buffer = BytesIO()
		workbook.save(buffer)
		buffer.seek(0)
		document = Document.objects.create(
			title='Dashboard Excel',
			slug='dashboard-excel',
			document_type='excel',
			summary='Resumen',
			file=SimpleUploadedFile('dashboard.xlsx', buffer.getvalue(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'),
			is_published=True,
		)

		data = dashboard_preview_data(document)
		self.assertEqual(data['annual'][1]['label'], '2022')
		self.assertEqual(data['annual'][1]['value'], '110.410')
		self.assertEqual(data['annual'][1]['detail'], '22,3%')
		self.assertEqual(data['annual'][1]['growth'], '+4,8%')
		self.assertEqual(data['kpis'][0]['value'], '495.272')
		self.assertEqual(data['kpis'][1]['value'], '2022')
		self.assertEqual(data['kpis'][2]['value'], '1.020')
		self.assertIn('MASCULINO', {item['label'] for item in data['gender']})
		self.assertEqual(data['ages'][0]['label'], 'ADOLESCENTES')
		self.assertEqual(
			next(item for item in data['regional_annual'] if item['region'] == 'REGION ANDINA' and item['year'] == '2022')['value'],
			'68.957',
		)

	def test_homepage_insights_computes_svg_and_metrics(self):
		statistics = {
			'annual': [{'label': '2021', 'value': '105.316'}, {'label': '2022', 'value': '110.410'}, {'label': '2023', 'value': '89.714'}],
			'regions': [{'label': 'REGION ANDINA', 'value': '315.074'}],
			'gender': [{'label': 'MASCULINO', 'value': '298.881'}],
			'ages': [{'label': 'ADULTOS', 'value': '455.211'}],
		}
		from .views import homepage_insights
		insights = homepage_insights(statistics)
		self.assertEqual(insights['peak_year'], '2022')
		self.assertEqual(len(insights['chart_points']), 3)
		self.assertTrue(insights['chart_polyline'])
		self.assertEqual(insights['chart_points'][1]['cx'], 160)
		self.assertEqual(insights['chart_points'][1]['cy'], 75)

	def test_homepage_summary_service_builds_insights(self):
		from .services import build_homepage_summary
		statistics = {
			'annual': [{'label': '2021', 'value': '105.316'}, {'label': '2022', 'value': '110.410'}, {'label': '2023', 'value': '89.714'}],
			'regions': [{'label': 'REGION ANDINA', 'value': '315.074'}],
			'gender': [{'label': 'MASCULINO', 'value': '298.881'}],
			'ages': [{'label': 'ADULTOS', 'value': '455.211'}],
		}
		insights = build_homepage_summary(statistics)
		self.assertEqual(insights['peak_year'], '2022')
		self.assertEqual(insights['top_region']['label'], 'REGION ANDINA')
		self.assertTrue(insights['chart_polyline'])

	def test_chat_answers_cases_vs_victims(self):
		response = self.client.post(reverse('portal:chat'), data={'category': 'dashboard', 'question': 'dashboard-cases-vs-victims'}, content_type='application/json')
		self.assertEqual(response.status_code, 200)
		self.assertIn('332.390', response.json()['answer'])
		self.assertIn('495.272', response.json()['answer'])



class DocumentValidationTests(TestCase):
	def test_rejects_invalid_file_extension(self):
		with self.assertRaises(Exception):
			Document.objects.create(
				title='Documento inválido',
				slug='documento-invalido',
				document_type='word',
				summary='No debería guardar',
				file=SimpleUploadedFile('invalido.txt', b'bad', content_type='text/plain'),
			)

	def test_rejects_invalid_pdf_extension(self):
		with self.assertRaises(Exception):
			Document(
				title='PDF inválido',
				slug='pdf-invalido',
				document_type='word',
				summary='No debería guardar',
				file=SimpleUploadedFile('file.docx', b'fake-doc', content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document'),
				pdf_file=SimpleUploadedFile('archivo.pdfx', b'bad', content_type='application/pdf'),
			).save()

	def test_rejects_file_larger_than_max_upload_size(self):
		with override_settings(MAX_UPLOAD_SIZE=10):
			with self.assertRaises(Exception):
				Document(
					title='Archivo grande',
					slug='archivo-grande',
					document_type='word',
					summary='No debería guardar',
					file=SimpleUploadedFile('file.docx', b'x' * 20, content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document'),
				).save().save()


class RepositoryDocumentSyncTests(TestCase):
	def test_sync_command_copies_and_publishes_repository_documents(self):
		with tempfile.TemporaryDirectory(prefix='repository-documents-test-') as directory:
			base_dir = Path(directory)
			source_dir = base_dir / 'documentos'
			source_dir.mkdir()
			media_root = base_dir / 'media'
			source_files = {
				'Documetacion_lesiones_personales.docx': b'docx content',
				'Documetacion_lesiones_personales.pdf': b'%PDF',
				'lesiones_personales_0 (1) (1).xlsx': b'xlsx content',
				'Presentacion De Proyecto.pptx': b'pptx content',
			}
			for filename, contents in source_files.items():
				(source_dir / filename).write_bytes(contents)

			with override_settings(BASE_DIR=base_dir, MEDIA_ROOT=media_root):
				management.call_command('sync_repository_documents', verbosity=0)
				management.call_command('sync_repository_documents', verbosity=0)

				self.assertEqual(Document.objects.filter(is_published=True).count(), 3)
				document = Document.objects.get(slug='lesiones-personales-en-colombia-20212025')
				self.assertEqual(document.file.name, 'documents/repository/lesiones_personales.xlsx')
				with document.file.open('rb') as saved_file:
					self.assertEqual(saved_file.read(), b'xlsx content')
				response = self.client.get(document.file.url)
				self.assertEqual(response.status_code, 200)
				self.assertEqual(b''.join(response.streaming_content), b'xlsx content')
				word_document = Document.objects.get(slug='documentacion-del-proyecto')
				self.assertTrue(default_storage.exists(word_document.pdf_file.name))
